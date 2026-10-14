# -*- coding: utf-8 -*-
"""
sheets_b.py — شیت‌های دادهای: اعضا و تراکنش‌ها.

شیت «تراکنش‌ها» قلب کارپوشه است: ستون‌های A..J ورودی کاربر و ستون‌های K..Z
زنجیرهٔ محاسبهٔ شفاف (مبنا → نرخ → کمیسیون → کف/سقف → مالیات → سهم‌ها → دوره).
همهٔ فرمول‌ها به نام‌های محدوده (تنظیمات/فهرست‌ها/تقویم) ارجاع می‌دهند، نه به عدد ثابت.
"""

from __future__ import annotations

from openpyxl.formatting.rule import CellIsRule, ColorScaleRule, DataBarRule, FormulaRule
from openpyxl.styles import Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table

import content as C
import theme as T
from jalali import from_jalali_text
from theme import L, aln, banner, brd_bottom, brd_box, fld, fnt, heights, merge, note_line, put, section, style_sheet, table_header, widths, zebra


# ═══════════════════════════════════════════════════ تعریف ستون‌های تراکنش

# (حرف، عنوان، نوع، عرض، قالب عدد) — نوع: input | calc | text
TX_COLUMNS = [
    ("A", "شمارهٔ پرونده", "text", 15, T.NF_TEXT),
    ("B", "تاریخ", "input", 12, T.NF_TEXT),
    ("C", "عضو", "input", 20, T.NF_TEXT),
    ("D", "نوع خدمت", "input", 22, T.NF_TEXT),
    ("E", "کارفرما", "input", 24, T.NF_TEXT),
    ("F", "مبلغ کارشناسی", "input", 18, T.NF_RIAL),
    ("G", "ضریب خدمت", "input", 10, T.NF_FACTOR),
    ("H", "سهم عضو بازنویسی", "input", 15, T.NF_PCT1),
    ("I", "وضعیت", "input", 15, T.NF_TEXT),
    ("J", "یادداشت", "input", 22, T.NF_TEXT),
    ("K", "مبلغ مبنا", "calc", 17, T.NF_RIAL),
    ("L", "نرخ پایه", "calc", 10, T.NF_PCT1),
    ("M", "نرخ مؤثر", "calc", 10, T.NF_PCT1),
    ("N", "کمیسیون خام", "calc", 17, T.NF_RIAL),
    ("O", "کمیسیون پس از کف و سقف", "calc", 19, T.NF_RIAL),
    ("P", "مالیات و عوارض", "calc", 15, T.NF_RIAL),
    ("Q", "جمع با مالیات", "calc", 17, T.NF_RIAL),
    ("R", "درصد نهایی سهم", "calc", 12, T.NF_PCT1),
    ("S", "مبلغ سهم عضو", "calc", 16, T.NF_RIAL),
    ("T", "سهم شرکت", "calc", 16, T.NF_RIAL),
    ("U", "سال شمسی", "calc", 9, T.NF_INT),
    ("V", "ماه شمسی", "calc", 9, T.NF_INT),
    ("W", "کلید دوره", "calc", 11, T.NF_TEXT),
    ("X", "فصل", "calc", 10, T.NF_TEXT),
    ("Y", "مهلت پرداخت", "calc", 12, T.NF_TEXT),
    ("Z", "سرسید سپری شده", "calc", 13, T.NF_TEXT),
]

COL = {c[0]: i + 1 for i, c in enumerate(TX_COLUMNS)}
COL_NAME = {c[1]: c[0] for c in TX_COLUMNS}


def _tx_formula(letter, r):
    """فرمول ستون محاسباتی برای ردیف r — زنجیرهٔ کامل و شفاف محاسبه."""
    f = {
        "K": f'=IF($F{r}="","",$F{r}*IF($G{r}="",1,$G{r}))',
        "L": f'=IFERROR(IF($D{r}="","",INDEX(t_Svc_Rate,MATCH($D{r},lst_Service,0))),"")',
        "M": f'=IF($L{r}="","",$L{r}*IF(p_TierMode="پلکانی",LOOKUP($K{r},t_Tier_From,t_Tier_Rate),1))',
        "N": f'=IF($K{r}="","",$K{r}*$M{r})',
        "O": f'=IF($N{r}="","",IF($I{r}="ابطال",0,MAX(p_Min,MIN(p_Max,$N{r}))))',
        "P": f'=IF($O{r}="","",ROUND($O{r}*p_Tax,0))',
        "Q": f'=IF($O{r}="","",$O{r}+$P{r})',
        "R": (f'=IF($A{r}="","",IF($H{r}<>"",$H{r},'
              f'IFERROR(INDEX(t_Member_Share,MATCH($C{r},t_Members,0)),p_DefaultShare)))'),
        "S": f'=IF(OR($Q{r}="",$R{r}=""),"",ROUND($Q{r}*$R{r},0))',
        "T": f'=IF(OR($Q{r}="",$S{r}=""),"",$Q{r}-$S{r})',
        "U": f'=IF($B{r}="","",IFERROR(INDEX(cal_JYear,MATCH($B{r},cal_Key,0)),""))',
        "V": f'=IF($B{r}="","",IFERROR(INDEX(cal_JMonth,MATCH($B{r},cal_Key,0)),""))',
        "W": f'=IF($B{r}="","",IFERROR(INDEX(cal_YM,MATCH($B{r},cal_Key,0)),""))',
        "X": f'=IF($B{r}="","",IFERROR(INDEX(cal_Season,MATCH($B{r},cal_Key,0)),""))',
        "Y": (f'=IF($B{r}="","",IFERROR(INDEX(cal_Key,MATCH('
              f'INDEX(cal_Serial,MATCH($B{r},cal_Key,0))+p_Due,cal_Serial,0)),""))'),
        "Z": (f'=IF($B{r}="","",IF(OR($I{r}="پرداخت‌شده",$I{r}="ابطال"),"—",'
              f'IFERROR(IF(INDEX(cal_Serial,MATCH($B{r},cal_Key,0))+p_Due<TODAY(),"بله","خیر"),'
              f'"تاریخ نامعتبر")))'),
    }
    return f.get(letter)


# ═══════════════════════════════════════════════════════════════ شیت اعضا

def build_members(wb):
    ws = wb.create_sheet(L.S_MEM)
    style_sheet(ws, gridlines=False, zoom=110, tab_color=L.TAB_COLORS[L.S_MEM],
                freeze="A6", print_area="A1:J18")
    widths(ws, {"A": 22, "B": 14, "C": 13, "D": 9, "E": 11, "F": 17,
                "G": 17, "H": 17, "I": 17, "J": 13})

    banner(ws, "اعضا و بازرسان فنی",
           "اطلاعات پایهٔ پرسنل + عملکرد زندهٔ هر نفر؛ ستون‌های E تا J خودکار محاسبه می‌شوند", 10)
    section(ws, L.MEM_SEC, 1, 10, "اطلاعات پایه و عملکرد تجمیعی از دفتر تراکنش‌ها")
    table_header(ws, L.MEM_HDR, 1,
                 ["نام و نام خانوادگی", "نقش", "سهم پیش‌فرض", "فعال",
                  "تعداد پرونده", "مبلغ کارشناسی", "کمیسیون خام",
                  "سهم عضو", "سهم شرکت", "آخرین فعالیت"], height=30)

    for i, (name, role, share, active) in enumerate(C.MEMBERS):
        r = L.MEM_FIRST + i
        put(ws, r, 1, name, f=fnt(10, True, T.TEXT), fl=fld(T.INPUT_BG), a=aln("right", indent=1))
        put(ws, r, 2, role, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("center"))
        put(ws, r, 3, share, f=fnt(10, True, T.NAVY), fl=fld(T.INPUT_BG), a=aln("center"), nf=T.NF_PCT1)
        put(ws, r, 4, active, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("center"))
        # ── عملکرد زنده (فقط رکوردهای غیرابطال)
        put(ws, r, 5, f'=COUNTIFS(tbl_TX[عضو],$A{r},tbl_TX[وضعیت],"<>ابطال")',
            f=fnt(10, True, T.TEAL), a=aln("center"), nf=T.NF_INT)
        put(ws, r, 6, f'=SUMIFS(tbl_TX[مبلغ کارشناسی],tbl_TX[عضو],$A{r},tbl_TX[وضعیت],"<>ابطال")',
            f=fnt(9.5), a=aln("center"), nf=T.NF_RIAL)
        put(ws, r, 7, f'=SUMIFS(tbl_TX[کمیسیون خام],tbl_TX[عضو],$A{r},tbl_TX[وضعیت],"<>ابطال")',
            f=fnt(9.5), a=aln("center"), nf=T.NF_RIAL)
        put(ws, r, 8, f'=SUMIFS(tbl_TX[مبلغ سهم عضو],tbl_TX[عضو],$A{r},tbl_TX[وضعیت],"<>ابطال")',
            f=fnt(9.5, True, T.OK_FG), a=aln("center"), nf=T.NF_RIAL)
        put(ws, r, 9, f'=SUMIFS(tbl_TX[سهم شرکت],tbl_TX[عضو],$A{r},tbl_TX[وضعیت],"<>ابطال")',
            f=fnt(9.5), a=aln("center"), nf=T.NF_RIAL)
        put(ws, r, 10, f'=IFERROR(LOOKUP(2,1/(tbl_TX[عضو]=$A{r}),tbl_TX[تاریخ]),"—")',
            f=fnt(9.5, False, T.MUTED), a=aln("center"))
        ws.row_dimensions[r].height = 21

    last = L.MEM_LAST
    tbl = Table(displayName="T_Members", ref=f"A{L.MEM_HDR}:J{last}")
    tbl.headerRowCount = 1
    tbl.totalsRowShown = False
    ws.add_table(tbl)
    zebra(ws, L.MEM_FIRST, last, 1, 10)

    # ── ردیف‌های یدکی: آمادهٔ تایپ عضو جدید (محدودهٔ پویا آن‌ها را خودکار می‌بیند)
    for r in range(L.MEM_LAST + 1, L.MEM_SPARE + 1):
        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.fill = fld(T.INPUT_BG)
            cell.font = fnt(9.5, False, T.MUTED, italic=True)
            cell.alignment = aln("center") if c != 1 else aln("right", indent=1)
            cell.border = brd_bottom()
            if c == 3:
                cell.number_format = T.NF_PCT1
        ws.cell(row=r, column=1).value = "— ردیف یدکی برای عضو جدید —" if r == L.MEM_LAST + 1 else None
        ws.row_dimensions[r].height = 20

    # ── ردیف جمع کل (دو ردیف پایین‌تر، بیرون از جدول تا افزودن عضو، جدول را گسترش دهد)
    tr = L.MEM_TOTAL
    put(ws, tr, 1, "جمع کل", f=fnt(10.5, True, T.WHITE), fl=fld(T.NAVY), a=aln("center"))
    for c in range(2, 11):
        put(ws, tr, c, None, fl=fld(T.NAVY))
    summary = {
        5: f'=SUM(E{L.MEM_FIRST}:E{last})',
        6: f'=SUM(F{L.MEM_FIRST}:F{last})',
        7: f'=SUM(G{L.MEM_FIRST}:G{last})',
        8: f'=SUM(H{L.MEM_FIRST}:H{last})',
        9: f'=SUM(I{L.MEM_FIRST}:I{last})',
    }
    for c, formula in summary.items():
        put(ws, tr, c, formula, f=fnt(10, True, T.WHITE), fl=fld(T.NAVY),
            a=aln("center"), nf=T.NF_RIAL if c != 5 else T.NF_INT)
    put(ws, tr, 10, f'=TEXT(TODAY(),"yyyy/mm/dd")', f=fnt(9, False, T.ACCENT),
        fl=fld(T.NAVY), a=aln("center"))
    ws.row_dimensions[tr].height = 24

    note_line(ws, tr + 1, 1, 10,
              "افزودن عضو جدید: نام و سهم را در یکی از ردیف‌های یدکی (۱۶ تا ۱۸) وارد کنید؛ "
              "محدودهٔ نام‌گذاری‌شدهٔ پویا آن را بلافاصله به کشویی‌های تراکنش‌ها و گزارش‌ها اضافه می‌کند.")
    note_line(ws, tr + 2, 1, 10,
              "سهم پیش‌فرض هر عضو مبنای محاسبه است، مگر آنکه برای پروندهٔ خاصی در ستون H شیت تراکنش‌ها بازنویسی شود.")

    # ── فرمت شرطی
    ws.conditional_formatting.add(
        f"D{L.MEM_FIRST}:D{last}",
        FormulaRule(formula=[f'$D{L.MEM_FIRST}="خیر"'],
                    fill=fld(T.BAD_BG), font=fnt(9.5, False, T.BAD_FG), stopIfTrue=False))
    ws.conditional_formatting.add(
        f"E{L.MEM_FIRST}:E{last}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color=T.TEAL,
                    showValue=True, minLength=None, maxLength=None))
    ws.conditional_formatting.add(
        f"H{L.MEM_FIRST}:H{last}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color=T.OK_FG,
                    showValue=True, minLength=None, maxLength=None))
    ws.conditional_formatting.add(
        f"A{L.MEM_FIRST}:J{last}",
        FormulaRule(formula=[f'AND($D{L.MEM_FIRST}="خیر",$A{L.MEM_FIRST}<>"")'],
                    font=fnt(10, False, T.MUTED, italic=True), stopIfTrue=False))
    return ws


# ═══════════════════════════════════════════════════════════════ شیت تراکنش‌ها

def build_transactions(wb):
    ws = wb.create_sheet(L.S_TX)
    style_sheet(ws, gridlines=False, zoom=95, tab_color=L.TAB_COLORS[L.S_TX],
                freeze="C6", orientation="landscape", title_rows="5:5",
                print_area=f"A1:Z{L.TX_LAST}")

    widths(ws, {c[0]: c[3] for c in TX_COLUMNS})
    heights(ws, {1: 34, 2: 20})

    banner(ws, "دفتر تراکنش‌ها — ثبت و محاسبهٔ زندهٔ کمیسیون",
           "ستون‌های آبی (A تا J) ورودی شما هستند؛ ستون‌های K تا Z خودکار محاسبه می‌شوند. "
           "یک ردیف خالی هم لازم نیست بسازید: کافی است در ردیف بعدی جدول تایپ کنید.", 26)

    note_line(ws, L.TX_NOTE, 1, 26,
              "راهنمای سریع: تاریخ را به شکل 1405/07/22 وارد کنید · شرح خدمات را از فهرست کشویی انتخاب کنید · "
              "مبلغ کارشناسی را بدون کاما بنویسید (کاما خودش نمایش داده می‌شود) · ستون H را فقط وقتی پر کنید که سهم این پرونده با سهم پیش‌فرض عضو تفاوت دارد.")

    # ── سرستون دو بخشی: ورودی / محاسبه
    input_cols = [c[0] for c in TX_COLUMNS if c[2] == "input"]
    calc_cols = [c[0] for c in TX_COLUMNS if c[2] == "calc"]
    merge(ws, L.TX_HDR - 1, COL[input_cols[0]], L.TX_HDR - 1, COL[input_cols[-1]])
    put(ws, L.TX_HDR - 1, COL[input_cols[0]], "▼ ورودی کاربر (آبی) — فقط این ستون‌ها را پر کنید",
        f=fnt(9, True, T.INFO_FG), fl=fld(T.ACCENT), a=aln("center"))
    for c in range(COL[input_cols[0]], COL[input_cols[-1]] + 1):
        put(ws, L.TX_HDR - 1, c, fl=fld(T.ACCENT))
    merge(ws, L.TX_HDR - 1, COL[calc_cols[0]], L.TX_HDR - 1, COL[calc_cols[-1]])
    put(ws, L.TX_HDR - 1, COL[calc_cols[0]], "▼ محاسبهٔ خودکار (فرمول زنده) — دست نزنید",
        f=fnt(9, True, T.TEAL), fl=fld(T.TEAL_SOFT), a=aln("center"))
    for c in range(COL[calc_cols[0]], COL[calc_cols[-1]] + 1):
        put(ws, L.TX_HDR - 1, c, fl=fld(T.TEAL_SOFT))
    ws.row_dimensions[L.TX_HDR - 1].height = 18

    for letter, header, kind, width, nf in TX_COLUMNS:
        put(ws, L.TX_HDR, COL[letter], header, f=fnt(9, True, T.WHITE),
            fl=fld(T.TEAL if kind == "calc" else T.NAVY),
            a=aln("center"), b=Border(bottom=Side(style="medium", color=T.GOLD)))
    ws.row_dimensions[L.TX_HDR].height = 34

    # ── ردیف‌های داده
    sample_dates = {}
    for i, (jdate, case, member, service, client, amount, factor, share, status) in enumerate(C.SAMPLE_TX):
        r = L.TX_FIRST + i
        put(ws, r, COL["A"], case, f=fnt(9.5, True, T.NAVY), fl=fld(T.INPUT_BG), a=aln("center"))
        put(ws, r, COL["B"], jdate, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("center"))
        put(ws, r, COL["C"], member, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("right", indent=1))
        put(ws, r, COL["D"], service, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("right", indent=1))
        put(ws, r, COL["E"], client, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("right", indent=1))
        put(ws, r, COL["F"], amount, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("center"), nf=T.NF_RIAL)
        put(ws, r, COL["G"], factor, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("center"), nf=T.NF_FACTOR)
        if share is not None:
            put(ws, r, COL["H"], share, f=fnt(9.5), fl=fld(T.INPUT_BG), a=aln("center"), nf=T.NF_PCT1)
        put(ws, r, COL["I"], status, f=fnt(9.5, True), fl=fld(T.INPUT_BG), a=aln("center"))
        put(ws, r, COL["J"], "", fl=fld(T.INPUT_BG), a=aln("right", indent=1))
        sample_dates[case] = from_jalali_text(jdate)

    NF_BY_COL = {c[0]: c[4] for c in TX_COLUMNS}
    for r in range(L.TX_FIRST, L.TX_LAST + 1):
        # ورودی‌ها: پس‌زمینهٔ آبی برای همهٔ ۳۰۰ ردیف (ردیف‌های خالی هم آمادهٔ تایپ‌اند)
        for letter in input_cols:
            cell = ws.cell(row=r, column=COL[letter])
            cell.fill = fld(T.INPUT_BG)
            cell.number_format = NF_BY_COL[letter]
            cell.font = fnt(9.5, letter == "A", T.NAVY if letter == "A" else T.TEXT)
            cell.alignment = aln("right", indent=1) if letter in ("C", "D", "E", "J") else aln("center")
        # محاسبه‌ها: فرمول زنده برای همهٔ ردیف‌ها (هر ردیف جدید آمادهٔ استفاده است)
        for letter in calc_cols:
            put(ws, r, COL[letter], _tx_formula(letter, r),
                f=fnt(9.5, letter == "Z"), a=aln("center"), nf=NF_BY_COL[letter])
        ws.row_dimensions[r].height = 20

    tbl = Table(displayName="tbl_TX", ref=f"A{L.TX_HDR}:Z{L.TX_LAST}")
    tbl.headerRowCount = 1
    tbl.totalsRowShown = False
    ws.add_table(tbl)

    # ═══════════ اعتبارسنجی داده
    def dv(**kw):
        v = T.mk_dv(**kw)
        ws.add_data_validation(v)
        return v

    d = dv(type="list", formula1="=t_Members", allow_blank=True, showDropDown=False,
           promptTitle="عضو مسئول", prompt="از فهرست شیت «اعضا» انتخاب کنید.",
           errorTitle="عضو نامعتبر", error="این نام در شیت «اعضا» وجود ندارد.", errorStyle="stop")
    d.add(f"C{L.TX_FIRST}:C{L.TX_LAST}")

    d = dv(type="list", formula1="=lst_Service", allow_blank=True, showDropDown=False,
           promptTitle="نوع خدمت", prompt="نرخ پایه از شیت «فهرست‌ها» خوانده می‌شود.",
           errorTitle="خدمت نامعتبر", error="این خدمت در فهرست خدمات نیست.", errorStyle="stop")
    d.add(f"D{L.TX_FIRST}:D{L.TX_LAST}")

    d = dv(type="list", formula1="=lst_Contractor", allow_blank=True, showDropDown=False,
           promptTitle="کارفرما", prompt="از فهرست کارفرمایان انتخاب کنید.",
           errorTitle="کارفرمای نامعتبر", error="این کارفرما در فهرست نیست.", errorStyle="stop")
    d.add(f"E{L.TX_FIRST}:E{L.TX_LAST}")

    d = dv(type="list", formula1="=lst_Status", allow_blank=True, showDropDown=False,
           promptTitle="وضعیت پرونده", prompt="وضعیت «ابطال» از جمع‌ها حذف می‌شود.",
           errorTitle="وضعیت نامعتبر", error="این وضعیت در فهرست نیست.", errorStyle="stop")
    d.add(f"I{L.TX_FIRST}:I{L.TX_LAST}")

    d = dv(type="decimal", operator="greaterThan", formula1="0", allow_blank=True,
           errorTitle="مبلغ نامعتبر", error="مبلغ کارشناسی باید مثبت باشد.", errorStyle="stop")
    d.add(f"F{L.TX_FIRST}:F{L.TX_LAST}")

    d = dv(type="decimal", operator="between", formula1="0.1", formula2="5", allow_blank=True,
           errorTitle="ضریب نامعتبر", error="ضریب خدمت باید بین ۰٫۱ و ۵ باشد.", errorStyle="stop")
    d.add(f"G{L.TX_FIRST}:G{L.TX_LAST}")

    d = dv(type="decimal", operator="between", formula1="0", formula2="1", allow_blank=True,
           errorTitle="درصد نامعتبر", error="درصد سهم باید بین ۰ و ۱ باشد (مثلاً ۰٫۲۵ برای ۲۵٪).",
           errorStyle="stop")
    d.add(f"H{L.TX_FIRST}:H{L.TX_LAST}")

    d = dv(type="custom", formula1=f"=COUNTIF(cal_Key,$B{L.TX_FIRST})=1", allow_blank=True,
           errorTitle="تاریخ نامعتبر", error="تاریخ را به شکل 1405/07/22 و از بازهٔ تقویم وارد کنید.",
           errorStyle="stop")
    d.add(f"B{L.TX_FIRST}:B{L.TX_LAST}")

    d = dv(type="custom", formula1=f"=COUNTIF($A${L.TX_FIRST}:$A${L.TX_LAST},$A{L.TX_FIRST})<=1",
           allow_blank=True, errorTitle="پروندهٔ تکراری",
           error="این شمارهٔ پرونده قبلاً ثبت شده است.", errorStyle="stop")
    d.add(f"A{L.TX_FIRST}:A{L.TX_LAST}")

    # ═══════════ فرمت شرطی
    rng = f"A{L.TX_FIRST}:Z{L.TX_LAST}"
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND($A{L.TX_FIRST}<>"",MOD(ROW(),2)=0)'],
                         fill=fld(T.SOFT), stopIfTrue=False))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'$I{L.TX_FIRST}="ابطال"'],
                         font=fnt(9.5, False, T.MUTED, strike=True, italic=True), stopIfTrue=False))
    ws.conditional_formatting.add(
        f"A{L.TX_FIRST}:A{L.TX_LAST}",
        FormulaRule(formula=[f'AND($A{L.TX_FIRST}<>"",COUNTIF($A${L.TX_FIRST}:$A${L.TX_LAST},$A{L.TX_FIRST})>1)'],
                    fill=fld(T.BAD_BG), font=fnt(9.5, True, T.BAD_FG), stopIfTrue=True))
    ws.conditional_formatting.add(
        rng, FormulaRule(formula=[f'AND($A{L.TX_FIRST}<>"",OR($C{L.TX_FIRST}="",$D{L.TX_FIRST}="",$F{L.TX_FIRST}=""))'],
                         fill=fld(T.WARN_BG), stopIfTrue=False))
    ws.conditional_formatting.add(
        f"F{L.TX_FIRST}:F{L.TX_LAST}",
        CellIsRule(operator="lessThan", formula=["0"], fill=fld(T.BAD_BG), font=fnt(9.5, True, T.BAD_FG)))
    ws.conditional_formatting.add(
        f"Q{L.TX_FIRST}:Q{L.TX_LAST}",
        DataBarRule(start_type="num", start_value=0, end_type="max", color=T.TEAL, showValue=True))
    ws.conditional_formatting.add(
        f"I{L.TX_FIRST}:I{L.TX_LAST}",
        FormulaRule(formula=[f'$I{L.TX_FIRST}="پرداخت‌شده"'], fill=fld(T.OK_BG),
                    font=fnt(9.5, True, T.OK_FG), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"I{L.TX_FIRST}:I{L.TX_LAST}",
        FormulaRule(formula=[f'OR($I{L.TX_FIRST}="معوق",$I{L.TX_FIRST}="در انتظار پرداخت")'],
                    fill=fld(T.WARN_BG), font=fnt(9.5, True, T.WARN_FG), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"I{L.TX_FIRST}:I{L.TX_LAST}",
        FormulaRule(formula=[f'$I{L.TX_FIRST}="ابطال"'], fill=fld(T.BAD_BG),
                    font=fnt(9.5, True, T.BAD_FG), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"Z{L.TX_FIRST}:Z{L.TX_LAST}",
        FormulaRule(formula=[f'$Z{L.TX_FIRST}="بله"'], fill=fld(T.BAD_BG),
                    font=fnt(9.5, True, T.BAD_FG), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"Z{L.TX_FIRST}:Z{L.TX_LAST}",
        FormulaRule(formula=[f'$Z{L.TX_FIRST}="تاریخ نامعتبر"'], fill=fld(T.WARN_BG),
                    font=fnt(9.5, False, T.WARN_FG), stopIfTrue=True))
    ws.conditional_formatting.add(
        f"M{L.TX_FIRST}:M{L.TX_LAST}",
        ColorScaleRule(start_type="min", start_color="FFFFFF", mid_type="percentile", mid_value=50,
                       mid_color="DCE9F5", end_type="max", end_color=T.TEAL))
    return ws
