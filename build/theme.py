# -*- coding: utf-8 -*-
"""
theme.py — سیستم طراحی (Design System) کارپوشهٔ «کمیسیون فنی».

تمام رنگ‌ها، فونت‌ها، قالب‌های عددی، ضخامت‌ها و مختصات ساختاری در همین فایل
متمرکز است تا ظاهر کارپوشه یکدست بماند و تغییر تم در آینده با ویرایش همین
فایل ممکن باشد. هیچ شیت‌سازی نباید رنگ یا فونت دلخواه داخل خود تعریف کند.
"""

from __future__ import annotations

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ═══════════════════════════════════════════════════════════════ پالت رنگ

NAVY = "12314F"        # سرمه‌ای اصلی — نوار عنوان، سرستون جداول
NAVY_DEEP = "0B2438"   # تیره‌تر — پانویس و نوارهای ثانویه
BLUE = "1F5C8B"        # آبی ساختاری — خط تأکید، زیرعنوان
TEAL = "0E8C8C"        # فیروزه‌ای تأکید — اعداد کلیدی، نمودار
TEAL_SOFT = "E3F2F2"   # فیروزه‌ای بسیار روشن — کارت‌های KPI
SOFT = "F5F8FB"        # پس‌زمینهٔ نرم
LINE = "D7E0E8"        # خط جداکنندهٔ hairline
INPUT_BG = "EAF2FD"    # پس‌زمینهٔ سلول‌های ورودی
CALC_BG = "FBFCFD"     # پس‌زمینهٔ بسیار روشن ستون‌های محاسباتی
ACCENT = "DCE9F5"      # آبی روشن برای نوارهای راهنما
WHITE = "FFFFFF"
TEXT = "1B2733"        # متن اصلی
MUTED = "64748B"       # متن ثانویه
GOLD = "C9A227"        # طلایی — تأکید ویژه

OK_FG, OK_BG = "1E8449", "E8F5EE"
WARN_FG, WARN_BG = "B7791F", "FDF3E3"
BAD_FG, BAD_BG = "B03A2E", "FBEAE8"
INFO_FG, INFO_BG = "1F5C8B", "EAF2FD"

# ═══════════════════════════════════════════════════════════════ فونت

FONT_NAME = "Tahoma"


def fnt(size=10, bold=False, color=TEXT, italic=False, name=FONT_NAME, underline=None, strike=False):
    """فونت با پیش‌فرض‌های تم."""
    return Font(name=name, size=size, bold=bold, color=color, italic=italic,
                underline=underline, strike=strike)


def fld(color):
    """پرکردن یکدست (solid)."""
    return PatternFill(fill_type="solid", start_color=color, end_color=color)


def _side(style="thin", color=LINE):
    return Side(style=style, color=color)


def brd_bottom(color=LINE, style="thin"):
    """فقط خط افقی پایین — زبان بصری اصلی این کارپوشه."""
    return Border(bottom=_side(style, color))


def brd_top(color=LINE, style="thin"):
    return Border(top=_side(style, color))


def brd_box(color=LINE, style="thin"):
    s = _side(style, color)
    return Border(left=s, right=s, top=s, bottom=s)


def brd_accent_bottom(color=BLUE, style="medium"):
    return Border(bottom=_side(style, color))


def mk_dv(**kw):
    """
    ساخت اعتبارسنجی داده به شکل استاندارد فایل اکسل:
    حذف «=» از ابتدای فرمول (قاعدهٔ XML) و فعال‌کردن پیام راهنما/خطا.
    """
    from openpyxl.worksheet.datavalidation import DataValidation
    for key in ("formula1", "formula2"):
        val = kw.get(key)
        if isinstance(val, str) and val.startswith("="):
            kw[key] = val[1:]
    kw.setdefault("showErrorMessage", True)
    kw.setdefault("showInputMessage", True)
    return DataValidation(**kw)


def aln(h="center", v="center", wrap=True, indent=0):
    """ترازبندی خواندنی برای شیت‌های راست‌به‌چپ (readingOrder=2)."""
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap, indent=indent, readingOrder=2)


ALN_RIGHT = aln("right")
ALN_LEFT = aln("left")
ALN_CENTER = aln("center")
ALN_NOWRAP_CENTER = aln("center", wrap=False)
ALN_NOWRAP_RIGHT = aln("right", wrap=False)

# ═══════════════════════════════════════════════════════════════ قالب اعداد

NF_RIAL = '#,##0" ریال"'
NF_RIAL_RED = '#,##0" ریال";[Red](#,##0)" ریال"'
NF_NUM = "#,##0"
NF_NUM2 = "#,##0.00"
NF_PCT1 = "0.0%"
NF_PCT0 = "0%"
NF_FACTOR = "0.000"
NF_INT = "0"
NF_DATE = "yyyy/mm/dd"
NF_DATE_FA = "[$-fa-IR,16]yyyy/mm/dd;@"
NF_TEXT = "@"
NF_QTY = '#,##0" مورد"'

# ═══════════════════════════════════════════════════════════════ مختصات ساختاری

class LAYOUT:
    """مختصات ثابت همهٔ شیت‌ها — تنها منبع حقیقت برای آدرس‌دهی فرمول‌ها."""

    # ── نام شیت‌ها
    S_HELP = "راهنما"
    S_SET = "تنظیمات"
    S_LISTS = "فهرست‌ها"
    S_MEM = "اعضا"
    S_TX = "تراکنش‌ها"
    S_CALC = "ماشین‌حساب"
    S_REP = "گزارش‌ها"
    S_DASH = "داشبورد"
    S_QC = "کنترل کیفیت"
    S_CAL = "تقویم"

    ORDER = [S_HELP, S_SET, S_LISTS, S_MEM, S_TX, S_CALC, S_REP, S_DASH, S_QC, S_CAL]
    TAB_COLORS = {
        S_HELP: GOLD, S_SET: NAVY_DEEP, S_LISTS: BLUE, S_MEM: BLUE,
        S_TX: TEAL, S_CALC: TEAL, S_REP: NAVY, S_DASH: NAVY, S_QC: BAD_FG,
        S_CAL: MUTED,
    }

    # ── تنظیمات: پارامترهای عمومی (مقدارها در ستون B)
    SET_SEC, SET_HDR = 4, 5
    SET_FIRST, SET_LAST = 6, 16               # ۱۱ پارامتر
    SET_COMPANY, SET_CURRENCY, SET_YEAR, SET_TAX = 6, 7, 8, 9
    SET_MIN, SET_MAX, SET_DUE, SET_MODE = 10, 11, 12, 13
    SET_SHARE, SET_LOWACT, SET_TARGET = 14, 15, 16
    SET_TIER_SEC, SET_TIER_HDR = 18, 19
    SET_TIER_FIRST, SET_TIER_LAST = 20, 23    # ۴ پله
    SET_NOTE_SEC, SET_NOTE_FIRST = 27, 28

    # ── فهرست‌ها (میان هر بلوک ۳ ردیف یدکی برای افزودن گزینهٔ جدید)
    SVC_SEC, SVC_HDR, SVC_FIRST, SVC_LAST, SVC_SPARE = 4, 5, 6, 13, 16
    STA_SEC, STA_HDR, STA_FIRST, STA_LAST, STA_SPARE = 17, 18, 19, 25, 28
    CLI_SEC, CLI_HDR, CLI_FIRST, CLI_LAST, CLI_SPARE = 29, 30, 31, 38, 41
    CAT_SEC, CAT_HDR, CAT_FIRST, CAT_LAST, CAT_SPARE = 42, 43, 44, 49, 52
    ROL_SEC, ROL_HDR, ROL_FIRST, ROL_LAST, ROL_SPARE = 53, 54, 55, 60, 63
    YN_SEC, YN_HDR, YN_FIRST, YN_LAST, YN_SPARE = 64, 65, 66, 68, 71
    MRG_SEC, MRG_HDR, MRG_FIRST, MRG_LAST, MRG_SPARE = 72, 73, 74, 75, 78

    # ── اعضا (۳ ردیف یدکی برای افزودن عضو جدید)
    MEM_SEC, MEM_HDR, MEM_FIRST, MEM_LAST = 4, 5, 6, 15       # ۱۰ عضو
    MEM_SPARE, MEM_TOTAL = 18, 20

    # ── تراکنش‌ها
    TX_NOTE = 3
    TX_HDR = 5
    TX_FIRST, TX_COUNT = 6, 300
    TX_LAST = TX_FIRST + TX_COUNT - 1                          # ۳۰۵
    TX_COLS = 26                                               # A..Z

    # ── تقویم
    CAL_HDR = 5
    CAL_FIRST = 6
    CAL_ROWS = 5114                                            # 2023-01-01 .. 2036-12-31
    CAL_LAST = CAL_FIRST + CAL_ROWS - 1                        # ۵۱۱۹

    # ── گزارش‌ها (چهار بلوک تمام‌عرض)
    REP_B1_SEC, REP_B1_HDR, REP_B1_FIRST = 5, 6, 7             # اعضا (۱۰ ردیف)
    REP_B1_LAST, REP_B1_TOTAL = 16, 17
    REP_B2_SEC, REP_B2_HDR, REP_B2_FIRST = 19, 20, 21          # ماه‌ها (۱۲ ردیف)
    REP_B2_LAST, REP_B2_TOTAL = 32, 33
    REP_B3_SEC, REP_B3_HDR, REP_B3_FIRST = 35, 36, 37          # خدمات (۸ ردیف)
    REP_B3_LAST, REP_B3_TOTAL = 44, 45
    REP_B4_SEC, REP_B4_HDR, REP_B4_FIRST = 47, 48, 49          # وضعیت‌ها (۷ ردیف)
    REP_B4_LAST, REP_B4_TOTAL = 55, 56
    REP_COLS = 10                                              # A..J

    # ── داشبورد
    DASH_FILTER_ROW = 3
    DASH_KPI_ROW = 5
    DASH_CHART_ROW = 10
    DASH_SPAN = 16                                             # A..P

    # ── کنترل کیفیت
    QC_SEC, QC_HDR, QC_FIRST = 4, 5, 6
    QC_TESTS = 12
    QC_LAST = QC_FIRST + QC_TESTS - 1


# ═══════════════════════════════════════════════════════════════ توابع کمکی نوشتن

def put(ws, r, c, v=None, *, f=None, fl=None, a=None, b=None, nf=None):
    """نوشتن یک سلول با استایل اختیاری و بازگرداندن خود سلول."""
    cell = ws.cell(row=r, column=c)
    if v is not None:
        cell.value = v
    if f is not None:
        cell.font = f
    if fl is not None:
        cell.fill = fl
    if a is not None:
        cell.alignment = a
    if b is not None:
        cell.border = b
    if nf is not None:
        cell.number_format = nf
    return cell


def put_row(ws, r, c0, values, *, f=None, fl=None, a=None, b=None, nf=None, skip_none=True):
    """نوشتن یک ردیف افقی از مقادیر."""
    for i, v in enumerate(values):
        if skip_none and v is None:
            continue
        put(ws, r, c0 + i, v, f=f, fl=fl, a=a, b=b, nf=nf)


def merge(ws, r1, c1, r2, c2):
    ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)


def widths(ws, spec):
    """spec: دیکشنری {حرف ستون: عرض} یا {شمارهٔ ستون: عرض}"""
    for key, w in spec.items():
        letter = key if isinstance(key, str) else get_column_letter(key)
        ws.column_dimensions[letter].width = w


def heights(ws, spec):
    for r, h in spec.items():
        ws.row_dimensions[r].height = h


def band(ws, r, c1, c2, *, height=None, fill=SOFT, border=None):
    """نوار افقی ادغام‌شده در یک ردیف."""
    if c2 > c1:
        merge(ws, r, c1, r, c2)
    for c in range(c1, c2 + 1):
        put(ws, r, c, fl=fill, border=border)
    if height:
        ws.row_dimensions[r].height = height


# ═══════════════════════════════════════════════════════════════ اجزای بصری استاندارد

def banner(ws, title, subtitle, span, *, row=1):
    """
    نوار عنوان شیت: نام شیت (درشت، سفید روی سرمه‌ای) + یک جملهٔ توضیح نقش شیت.
    `span` = آخرین شمارهٔ ستون محدودهٔ نوار.
    """
    merge(ws, row, 1, row, span)
    put(ws, row, 1, title,
        f=fnt(17, True, WHITE), fl=fld(NAVY), a=aln("right", indent=1))
    for c in range(1, span + 1):
        put(ws, row, c, fl=fld(NAVY))
    ws.row_dimensions[row].height = 36

    merge(ws, row + 1, 1, row + 1, span)
    put(ws, row + 1, 1, subtitle,
        f=fnt(9.5, False, WHITE), fl=fld(BLUE), a=aln("right", indent=1))
    for c in range(1, span + 1):
        put(ws, row + 1, c, fl=fld(BLUE))
    ws.row_dimensions[row + 1].height = 22

    # خط طلایی زیر نوار عنوان — امضای بصری کارپوشه
    for c in range(1, span + 1):
        ws.cell(row=row + 1, column=c).border = brd_bottom(GOLD, "medium")
    return row + 2


def section(ws, r, c1, c2, text, *, fill=None, color=NAVY):
    """عنوان یک بلوک/گزارش با خط تأکید فیروزه‌ای."""
    if c2 > c1:
        merge(ws, r, c1, r, c2)
    for c in range(c1, c2 + 1):
        put(ws, r, c, fl=fill)
    put(ws, r, c1, text, f=fnt(11.5, True, color), a=aln("right", indent=1))
    for c in range(c1, c2 + 1):
        cell = ws.cell(row=r, column=c)
        cell.border = Border(bottom=Side(style="medium", color=TEAL))
    ws.row_dimensions[r].height = 26
    return r + 1


def table_header(ws, r, c1, values, *, fill=NAVY, height=30, color=WHITE, size=9.5):
    """سرستون جدول: سرمه‌ای، سفید، وسط‌چین، خط تأکید آبی در پایین."""
    for i, v in enumerate(values):
        put(ws, r, c1 + i, v,
            f=fnt(size, True, color), fl=fld(fill), a=aln("center"),
            b=Border(bottom=Side(style="medium", color=BLUE)))
    ws.row_dimensions[r].height = height
    return r + 1


def zebra(ws, r1, r2, c1, c2, *, base=WHITE, alt=SOFT, border=True):
    """راه‌راه کردن ردیف‌های یک محدوده به سبک مینیمال (فقط خط افقی)."""
    for r in range(r1, r2 + 1):
        fl = fld(alt) if (r - r1) % 2 else fld(base)
        for c in range(c1, c2 + 1):
            cell = ws.cell(row=r, column=c)
            cell.fill = fl
            if border:
                cell.border = brd_bottom(LINE)


def kpi_card(ws, r, c1, c2, label, value, *, fmt=NF_RIAL, fill=TEAL_SOFT, fg=NAVY, note=None):
    """کارت شاخص کلیدی: برچسب ریز بالا، عدد بزرگ پایین."""
    merge(ws, r, c1, r, c2)
    put(ws, r, c1, label, f=fnt(9, True, MUTED), fl=fld(fill), a=aln("center"))
    for c in range(c1, c2 + 1):
        put(ws, r, c, fl=fld(fill))
    ws.row_dimensions[r].height = 20

    merge(ws, r + 1, c1, r + 1, c2)
    put(ws, r + 1, c1, value, f=fnt(13, True, fg), fl=fld(fill), a=aln("center"), nf=fmt)
    for c in range(c1, c2 + 1):
        put(ws, r + 1, c, fl=fld(fill), b=Border(bottom=Side(style="medium", color=fg)))
    ws.row_dimensions[r + 1].height = 28

    if note is not None:
        merge(ws, r + 2, c1, r + 2, c2)
        put(ws, r + 2, c1, note, f=fnt(8, False, MUTED), fl=fld(fill), a=aln("center"))
        for c in range(c1, c2 + 1):
            put(ws, r + 2, c, fl=fld(fill))
        ws.row_dimensions[r + 2].height = 15


def note_line(ws, r, c1, c2, text, *, fill=ACCENT, color=INFO_FG, size=9):
    """نوار یادداشت/راهنما با نوار رنگی سمت راست."""
    merge(ws, r, c1, r, c2)
    put(ws, r, c1, text, f=fnt(size, False, color), fl=fld(fill), a=aln("right", indent=1))
    for c in range(c1, c2 + 1):
        put(ws, r, c, fl=fld(fill))
    ws.cell(row=r, column=c2).border = Border(left=Side(style="thick", color=TEAL))
    ws.row_dimensions[r].height = 20


def nav_link(ws, r, c, target_sheet, text=None):
    """لینک ناوبری داخلی به شیت دیگر (کارت کوچک)."""
    cell = put(ws, r, c, text or ("→ " + target_sheet),
               f=fnt(9, True, WHITE), fl=fld(BLUE), a=aln("center"))
    cell.hyperlink = f"#'{target_sheet}'!A1"
    return cell


def style_sheet(ws, *, rtl=True, gridlines=False, zoom=100, tab_color=None,
                freeze=None, orientation="portrait", fit_width=1, fit_height=0,
                title_rows=None, title_cols=None, footer=True, print_area=None):
    """اعمال تنظیمات ساختاری و چاپی یکدست روی شیت."""
    ws.sheet_view.rightToLeft = rtl
    ws.sheet_view.showGridLines = gridlines
    ws.sheet_view.zoomScale = zoom
    if tab_color:
        ws.sheet_properties.tabColor = tab_color
    if freeze:
        ws.freeze_panes = freeze

    ws.page_setup.orientation = orientation
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = fit_width
    ws.page_setup.fitToHeight = fit_height
    ws.print_options.horizontalCentered = True
    if title_rows:
        ws.print_title_rows = title_rows
    if title_cols:
        ws.print_title_cols = title_cols
    if print_area:
        ws.print_area = print_area
    if footer:
        ws.oddFooter.center.text = "&A — صفحهٔ &P از &N"
        ws.oddFooter.center.size = 8
        ws.oddFooter.center.color = MUTED
        ws.evenFooter.center.text = ws.oddFooter.center.text


# نام کوتاه برای دسترسی راحت همهٔ ماژول‌ها
L = LAYOUT
